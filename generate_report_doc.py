import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Set page margins (1 inch = 1440 dxa)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("BÁO CÁO BÀI THI THỰC HÀNH - MMA301")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Navy

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    sub_run = sub_p.add_run("ỨNG DỤNG DI ĐỘNG BOOKSHELF (EXPO / REACT NATIVE & FIREBASE)")
    sub_run.font.size = Pt(13)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Royal Blue

    # Section 1: Template Thông tin sinh viên (Copy đúng mẫu đề bài yêu cầu)
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1_run = h1.add_run("1. THÔNG TIN BÀI NỘP (STUDENT INFORMATION)")
    h1_run.font.size = Pt(14)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    desc_info = doc.add_paragraph("Mẫu thông tin bắt buộc theo định dạng đề bài yêu cầu (giữ nguyên các tiêu đề và định dạng URL có https://):")
    desc_info.paragraph_format.space_after = Pt(6)

    # Info Table
    table = doc.add_table(rows=7, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    info_data = [
        ("Student Code:", "QE200012 (Thay đổi mã số sinh viên của bạn)"),
        ("Full Name:", "Nguyen Van A (Thay đổi họ tên đầy đủ của bạn)"),
        ("GitHub Repository URL:", "https://github.com/DevTurtle-Hub/PE_MMA"),
        ("Expo Project Link:", "https://expo.dev/accounts/username/projects/bookshelf/builds/xxxxxxxx"),
        ("Firebase Project ID:", "pemma301"),
        ("Test account email:", "grader.test@example.com"),
        ("Test account password:", "Password123!"),
    ]

    for idx, (label, val) in enumerate(info_data):
        row = table.rows[idx]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]
        
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.5)
        
        set_cell_background(cell_lbl, 'F1F5F9')
        set_cell_margins(cell_lbl, 100, 100, 140, 140)
        set_cell_margins(cell_val, 100, 100, 140, 140)
        
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(0)
        r_lbl = p_lbl.add_run(label)
        r_lbl.bold = True
        r_lbl.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(0)
        r_val = p_val.add_run(val)
        if "https://" in val:
            r_val.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            r_val.underline = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 2: Danh sách 16 ảnh chụp màn hình theo đúng thứ tự đề bài
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    h2_run = h2.add_run("2. CHI TIẾT 16 ẢNH CHỤP MÀN HÌNH MINH CHỨNG (SCREENSHOTS)")
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    intro_sc = doc.add_paragraph(
        "Theo yêu cầu của đề thi (Mục 3), mỗi tính năng/màn hình cần có 1 ảnh chụp rõ ràng từ thiết bị thật hoặc máy ảo "
        "đã đăng nhập tài khoản test (không dùng ảnh mockup thiết kế). Dưới đây là mô tả chi tiết giải pháp kỹ thuật, "
        "hướng dẫn thao tác lấy ảnh và vị trí chèn ảnh đúng chuẩn."
    )
    intro_sc.paragraph_format.space_after = Pt(12)

    screenshots = [
        {
            "id": 1,
            "title": "Ảnh 1: Firebase console — Authentication",
            "caption": "Firebase console — Authentication: the Users list showing your test account.",
            "purpose": "Chứng minh hệ thống Authentication đã kích hoạt Email/Password Provider và tài khoản test dành cho giám khảo đã được tạo sẵn trên hệ thống Firebase Cloud.",
            "technical": (
                "• Firebase Auth SDK: Sử dụng hàm createUserWithEmailAndPassword và signInWithEmailAndPassword từ firebase/auth.\n"
                "• Tài khoản thử nghiệm đã tạo: grader.test@example.com (UID duy nhất sinh bởi Firebase).\n"
                "• Thiết lập Security: Bảo đảm xác thực cấp hệ sinh thái Firebase, token phiên đăng nhập được quản lý tự động."
            ),
            "instruction": "Mở trình duyệt vào Firebase Console -> Project 'pemma301' -> Mục 'Build' -> 'Authentication' -> Tab 'Users'. Chụp toàn màn hình hiển thị dòng chứa email tài khoản test (grader.test@example.com)."
        },
        {
            "id": 2,
            "title": "Ảnh 2: Firebase console — Firestore data",
            "caption": "Firebase console — Firestore data: the books collection with at least 3 sample documents belonging to your test account.",
            "purpose": "Chứng minh cấu trúc cơ sở dữ liệu Cloud Firestore với collection duy nhất 'books', chứa ít nhất 3 documents mẫu thuộc quyền sở hữu của tài khoản test.",
            "technical": (
                "• Cấu trúc Collection: 'books' lưu trữ toàn bộ sách của người dùng.\n"
                "• Schema chuẩn hóa: id (document id), userId (UID chủ sở hữu), title (string <= 100), author (string <= 60), genre (Fiction/Non-Fiction/Academic/Other), status (To Read/Reading/Finished), rating (0-5 number), favorite (boolean), notes (string <= 300), createdAt (Timestamp).\n"
                "• Đảm bảo có ít nhất 3 document có trường userId khớp với UID của tài khoản test."
            ),
            "instruction": "Vào Firebase Console -> 'Firestore Database' -> Chọn collection 'books'. Chọn xem các document mẫu sao cho thấy rõ các trường (title, author, genre, status, userId). Chụp màn hình bảng dữ liệu."
        },
        {
            "id": 3,
            "title": "Ảnh 3: Firebase console — Security Rules",
            "caption": "Firebase console — Security Rules: the Rules tab showing the rules you actually have published, not just saved in an editor.",
            "purpose": "Chứng minh quy tắc bảo mật Cloud Firestore đã được Publish trực tiếp lên Firebase Cloud, từ chối mọi truy vấn bất hợp pháp hoặc dữ liệu sai từ backend.",
            "technical": (
                "• Quyền hạn: Chỉ người dùng đã đăng nhập (request.auth != null) mới có quyền thao tác.\n"
                "• Cô lập dữ liệu: Người dùng chỉ được đọc, tạo, cập nhật, xóa các tài liệu có userId == request.auth.uid.\n"
                "• Backend Validation: Kiểm tra bắt buộc title, author không rỗng, giới hạn độ dài ký tự, giới hạn enum genre & status, khoảng giá trị rating (0-5) và notes <= 300 ký tự."
            ),
            "instruction": "Vào Firebase Console -> 'Firestore Database' -> Tab 'Rules'. Chụp toàn bộ giao diện editor chứa mã nguồn file firestore.rules và thấy rõ nút 'Publish' đã chuyển trạng thái xuất bản thành công."
        },
        {
            "id": 4,
            "title": "Ảnh 4: Register screen — validation",
            "caption": "Register screen — validation: the form submitted empty, showing the error messages.",
            "purpose": "Chứng minh tính năng kiểm tra dữ liệu đầu vào của màn hình Đăng ký, hiển thị thông báo lỗi chuẩn xác khi bấm Submit form trống.",
            "technical": (
                "• File xử lý: src/screens/RegisterScreen.tsx kết hợp hàm validateRegisterForm tại src/utils/validators.ts.\n"
                "• Quy tắc hiển thị lỗi: Không hiện lỗi trước lần submit đầu tiên. Khi bấm Register với form trống, hiển thị đồng thời:\n"
                "  - Email trống: 'Email is required'\n"
                "  - Password trống: 'Password is required'\n"
                "  - Confirm password trống: 'Confirm password is required'\n"
                "• Cơ chế tự động: Khi người dùng gõ nhập hợp lệ vào trường nào, thông báo lỗi trường đó lập tức biến mất."
            ),
            "instruction": "Mở ứng dụng ở màn hình Home -> Bấm 'Register' -> Để trống cả 3 ô (Email, Password, Confirm Password) -> Bấm nút 'Register'. Chụp màn hình hiển thị cả 3 câu thông báo lỗi màu đỏ."
        },
        {
            "id": 5,
            "title": "Ảnh 5: Register screen — success",
            "caption": "Register screen — success: right after a successful sign-up, showing Registration successful.",
            "purpose": "Chứng minh quá trình đăng ký tài khoản mới thành công lên Firebase Authentication và phản hồi giao diện tức thì.",
            "technical": (
                "• File xử lý: src/screens/RegisterScreen.tsx & src/services/auth.service.ts.\n"
                "• Xử lý thành công: Gọi authService.registerUser(), tạo tài khoản thực trên Firebase Auth.\n"
                "• testID chuẩn: testID='form-success' hiển thị văn bản 'Registration successful'.\n"
                "• Điều hướng: Sau khi báo thành công, ứng dụng tự động chuyển người dùng vào hệ thống mà không bị giật lag."
            ),
            "instruction": "Tại màn hình Register, nhập một email mới hợp lệ (ví dụ: newuser@example.com), password >= 6 ký tự và confirm password khớp nhau. Bấm 'Register'. Chụp màn hình ngay khi hộp màu xanh 'Registration successful' xuất hiện."
        },
        {
            "id": 6,
            "title": "Ảnh 6: Login screen",
            "caption": "Login screen: the empty sign-in form.",
            "purpose": "Chứng minh giao diện màn hình Đăng nhập sạch sẽ, chuyên nghiệp, không hiển thị bất kỳ lỗi nào khi người dùng mới mở màn hình.",
            "technical": (
                "• File xử lý: src/screens/LoginScreen.tsx.\n"
                "• Giao diện: Thiết kế hiện đại với minh họa đồ họa 3D đọc sách, 2 ô nhập liệu Email và Password, nút 'Sign In' và liên kết chuyển nhanh sang Register.\n"
                "• Trạng thái ban đầu: hasSubmitted = false, errors = {}, không hiển thị bất kỳ cảnh báo lỗi nào."
            ),
            "instruction": "Từ màn hình Home, bấm nút 'Login' để chuyển sang màn hình Đăng nhập. Chụp màn hình khi form đang để trống hoàn toàn và chưa bấm đăng nhập."
        },
        {
            "id": 7,
            "title": "Ảnh 7: List screen — after login",
            "caption": "List screen — after login: showing the user’s e-mail, the Logout button, and at least 3 books already in the list.",
            "purpose": "Chứng minh màn hình danh sách sách (List Screen) sau khi đăng nhập thành công, hiển thị email tài khoản, nút Logout và danh sách ít nhất 3 cuốn sách.",
            "technical": (
                "• File xử lý: src/screens/ListScreen.tsx & src/contexts/AuthContext.tsx.\n"
                "• testID bắt buộc: testID='list-screen'.\n"
                "• Header: Hiển thị email người dùng hiện tại (currentUser.email) và nút Logout viền đỏ nổi bật.\n"
                "• Danh sách: FlatList hiển thị các cuốn sách được tải theo thời gian thực (onSnapshot listener) từ collection 'books' thuộc về userId này."
            ),
            "instruction": "Đăng nhập bằng tài khoản test (grader.test@example.com). Chụp màn hình danh sách sách sao cho thấy rõ: thanh header có email, nút Logout, nút [+ Add New Book] và ít nhất 3 cuốn sách có sẵn trong danh sách."
        },
        {
            "id": 8,
            "title": "Ảnh 8: Add Book form — validation",
            "caption": "Add Book form — validation: the form submitted empty or with bad input, showing the error messages.",
            "purpose": "Chứng minh biểu mẫu thêm sách kiểm tra hợp lệ toàn bộ các trường theo đúng thông báo lỗi quy định trong đề thi.",
            "technical": (
                "• File xử lý: src/components/books/BookForm.tsx & src/utils/validators.ts.\n"
                "• testID: testID='book-form'.\n"
                "• Các thông báo lỗi xuất hiện khi form rỗng hoặc nhập sai:\n"
                "  - Title: 'Title is required'\n"
                "  - Author: 'Author is required'\n"
                "  - Genre: 'Please choose a genre'\n"
                "  - Status: 'Please choose a status'\n"
                "  - Rating: 'Rating must be a whole number from 0 to 5' (nếu nhập sai)\n"
                "  - Notes: 'Notes must be at most 300 characters' (nếu vượt quá)"
            ),
            "instruction": "Tại List screen, bấm nút '[+ Add New Book]' để mở form thêm sách. Để trống toàn bộ các trường và bấm nút 'Add Book'. Chụp màn hình hiển thị form kèm các câu thông báo lỗi đỏ dưới từng ô nhập liệu."
        },
        {
            "id": 9,
            "title": "Ảnh 9: Add Book form — success",
            "caption": "Add Book form — success: right after adding a book, showing Book added and the new book now in the list.",
            "purpose": "Chứng minh thêm sách thành công: hiển thị thông báo 'Book added', form được xóa trắng và sách mới xuất hiện ngay lập tức trong danh sách không cần reload.",
            "technical": (
                "• File xử lý: src/components/books/BookForm.tsx & src/contexts/BooksContext.tsx.\n"
                "• testID thành công: testID='book-success' hiển thị văn bản 'Book added'.\n"
                "• Lưu trữ: Gọi createBookDocument() lưu lên Firestore, đồng thời Optimistic Update đưa sách vào state books của React Context ngay tức thì.\n"
                "• Sắp xếp: Sách mới được tự động xếp vào đúng vị trí bảng chữ cái (A-Z) theo tiêu đề."
            ),
            "instruction": "Nhập dữ liệu hợp lệ cho sách mới (ví dụ: Title: 'Clean Architecture', Author: 'Robert C. Martin', Genre: 'Academic', Status: 'Reading', Rating: 5). Bấm 'Add Book'. Chụp màn hình ngay khi hiện thông báo xanh 'Book added' và sách mới xuất hiện ở danh sách."
        },
        {
            "id": 10,
            "title": "Ảnh 10: Filters and search",
            "caption": "Filters and search: the genre filter, the status filter, the favorites-only switch, and the search box each changing what is shown.",
            "purpose": "Chứng minh hệ thống tìm kiếm và lọc hoạt động tức thời, kết hợp đồng thời (AND logic) giữa 4 điều kiện và nút Clear Filters.",
            "technical": (
                "• File xử lý: src/components/books/BookFilter.tsx & src/hooks/useBookFilters.ts.\n"
                "• Search Box: Tìm kiếm không phân biệt chữ hoa/thường theo Title hoặc Author.\n"
                "• Dropdown Genre: Lọc theo All hoặc từng thể loại (Fiction, Non-Fiction, Academic, Other).\n"
                "• Dropdown Status: Lọc theo All hoặc từng trạng thái (To Read, Reading, Finished).\n"
                "• Switch Favorites only: Lọc chỉ sách có favorite == true.\n"
                "• Nút Clear Filters: Reset toàn bộ 4 điều kiện về ban đầu chỉ với 1 lần chạm."
            ),
            "instruction": "Nhập một từ khóa vào ô tìm kiếm (hoặc chọn 1 Genre/Status cụ thể, hoặc gạt bật 'Favorites only') để danh sách lọc bớt các sách không khớp. Chụp màn hình thể hiện bộ lọc đang hoạt động và danh sách kết quả tương ứng."
        },
        {
            "id": 11,
            "title": "Ảnh 11: Summary",
            "caption": "Summary: book-count, favorite-count and the status breakdown visible, with more than one status represented.",
            "purpose": "Chứng minh bảng thống kê tổng hợp số lượng sách, số sách yêu thích và phân rã theo trạng thái đang hiển thị trên màn hình.",
            "technical": (
                "• File xử lý: src/components/books/BookSummary.tsx & src/utils/sorters.ts.\n"
                "• testID: testID='book-count' hiển thị tổng số sách (số nguyên thuần), testID='favorite-count' hiển thị số sách yêu thích.\n"
                "• Logic Status Breakdown: Đếm số sách theo từng trạng thái, sắp xếp theo trạng thái có nhiều sách nhất lên trước (nếu bằng nhau ưu tiên To Read -> Reading -> Finished).\n"
                "• Tính động: Tự động tính toán lại ngay khi bộ lọc hoặc tìm kiếm thay đổi."
            ),
            "instruction": "Chụp cận cảnh hoặc toàn màn hình khối Summary (Current View Summary) hiển thị rõ: ô 'Total Books', ô 'Favorites' và bảng 'Status Breakdown' có ít nhất 2 trạng thái khác nhau (ví dụ: Reading: 2, To Read: 1)."
        },
        {
            "id": 12,
            "title": "Ảnh 12: Detail screen — view mode",
            "caption": "Detail screen — view mode: one book’s full details (title, author, genre, status, rating, notes).",
            "purpose": "Chứng minh màn hình xem chi tiết sách hiển thị toàn bộ các trường dữ liệu của một cuốn sách được chọn từ danh sách.",
            "technical": (
                "• File xử lý: src/screens/DetailScreen.tsx.\n"
                "• Thông tin hiển thị: Title cỡ lớn, Author, Genre badge, Status badge, Rating (ví dụ '5 / 5'), phần Notes đầy đủ nội dung.\n"
                "• Các nút điều khiển: Nút ngôi sao yêu thích trên Header, nút 'Edit' và nút 'Delete'."
            ),
            "instruction": "Từ danh sách sách, bấm vào 1 cuốn sách bất kỳ có đầy đủ rating và notes để chuyển sang màn hình Detail. Chụp màn hình chế độ xem (View Mode) thể hiện đầy đủ các trường thông tin."
        },
        {
            "id": 13,
            "title": "Ảnh 13: Detail screen — edit mode, validation",
            "caption": "Detail screen — edit mode, validation: an edit submitted with bad input, showing the error messages.",
            "purpose": "Chứng minh chế độ chỉnh sửa (Edit Mode) của màn hình Detail kiểm tra tính hợp lệ của dữ liệu trước khi cho phép lưu.",
            "technical": (
                "• File xử lý: src/screens/DetailScreen.tsx.\n"
                "• Cơ chế chuyển đổi: Bấm nút 'Edit' biến các trường hiển thị thành TextInput và Picker kèm 2 nút 'Save' và 'Cancel'.\n"
                "• Validation: Sử dụng cùng hàm validateBookData. Nếu xóa trống Title/Author hoặc nhập Rating sai, thông báo lỗi xuất hiện ngay dưới trường đó và chặn thao tác ghi vào Firestore."
            ),
            "instruction": "Tại màn hình Detail, bấm nút 'Edit'. Xóa trắng ô Title hoặc Author rồi bấm nút 'Save'. Chụp ảnh màn hình hiển thị form chỉnh sửa kèm thông báo lỗi màu đỏ."
        },
        {
            "id": 14,
            "title": "Ảnh 14: Detail screen — after Save",
            "caption": "Detail screen — after Save: showing the edited values now saved.",
            "purpose": "Chứng minh lưu chỉnh sửa sách thành công: dữ liệu cập nhật trên Firestore và quay lại View Mode với giá trị mới.",
            "technical": (
                "• File xử lý: src/screens/DetailScreen.tsx & src/services/books.service.ts.\n"
                "• Quy trình: Gọi updateBookDocument(), cập nhật Firestore và đồng bộ state context, sau đó đóng chế độ Edit (setIsEditing(false)) trở về View Mode.\n"
                "• Chống trùng: Nếu chỉnh sửa tiêu đề và tác giả trùng với sách khác của người dùng, hệ thống sẽ từ chối lưu và báo lỗi."
            ),
            "instruction": "Sau khi nhập thông tin chỉnh sửa hợp lệ (ví dụ đổi Status từ 'To Read' sang 'Reading' hoặc thêm nội dung vào Notes), bấm nút 'Save'. Chụp ảnh màn hình View Mode hiển thị giá trị mới vừa được cập nhật."
        },
        {
            "id": 15,
            "title": "Ảnh 15: Favorite toggle and Delete",
            "caption": "Favorite toggle and Delete: a before/after pair showing the star toggled, and a book deleted and gone from the list.",
            "purpose": "Chứng minh thao tác bật/tắt yêu thích (Favorite Toggle) và thao tác xóa sách (Delete) có hộp thoại xác nhận, sách biến mất lập tức khỏi danh sách.",
            "technical": (
                "• File xử lý: src/screens/DetailScreen.tsx, src/components/books/BookItem.tsx, src/services/books.service.ts.\n"
                "• Favorite: Bấm ngôi sao gọi toggleFavoriteDocument(), icon đổi màu vàng/xám tức thời.\n"
                "• Delete: Bấm 'Delete' kích hoạt Alert.alert() xác nhận 'Confirm Delete'. Khi xác nhận, gọi deleteBookDocument(), xóa khỏi Firestore, tự động điều hướng navigation.goBack() và sách bị loại khỏi danh sách ngay lập tức."
            ),
            "instruction": "Ghép hoặc chèn 2 ảnh cạnh nhau: Ảnh A thể hiện bấm ngôi sao yêu thích (ngôi sao chuyển sang màu vàng) hoặc popup xác nhận xóa; Ảnh B thể hiện danh sách sau khi sách đã bị xóa thành công."
        },
        {
            "id": 16,
            "title": "Ảnh 16: Empty state",
            "caption": "Empty state: the books-empty screen with the text No books found.",
            "purpose": "Chứng minh trạng thái rỗng khi danh sách không có sách (tài khoản mới chưa có sách hoặc tìm kiếm không khớp kết quả), hiển thị component chuẩn theo yêu cầu.",
            "technical": (
                "• File xử lý: src/components/books/BookEmpty.tsx & src/screens/ListScreen.tsx.\n"
                "• testID bắt buộc: testID='books-empty' hiển thị nội dung chính xác 'No books found'.\n"
                "• Cấu hình FlatList: ListEmptyComponent={BookEmpty} tự động kích hoạt khi mảng dữ liệu filteredBooks rỗng."
            ),
            "instruction": "Gõ một từ khóa không tồn tại vào ô Search (ví dụ: 'xyz999') hoặc đăng nhập bằng một tài khoản mới tinh chưa thêm sách. Chụp ảnh màn hình hiển thị thẻ Empty State có biểu tượng sách và dòng chữ 'No books found'."
        },
    ]

    for item in screenshots:
        # Title of screenshot
        p_item_title = doc.add_paragraph()
        p_item_title.paragraph_format.space_before = Pt(14)
        p_item_title.paragraph_format.space_after = Pt(4)
        r_item_title = p_item_title.add_run(f"2.{item['id']} {item['title']}")
        r_item_title.font.size = Pt(12)
        r_item_title.font.bold = True
        r_item_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        # Purpose
        p_purp = doc.add_paragraph()
        p_purp.paragraph_format.space_after = Pt(2)
        r_purp_lbl = p_purp.add_run("• Mục đích kiểm tra: ")
        r_purp_lbl.bold = True
        p_purp.add_run(item['purpose'])

        # Technical details
        p_tech = doc.add_paragraph()
        p_tech.paragraph_format.space_after = Pt(2)
        r_tech_lbl = p_tech.add_run("• Triển khai kỹ thuật: \n")
        r_tech_lbl.bold = True
        p_tech.add_run(item['technical'])

        # How to capture
        p_inst = doc.add_paragraph()
        p_inst.paragraph_format.space_after = Pt(6)
        r_inst_lbl = p_inst.add_run("• Hướng dẫn chụp ảnh: ")
        r_inst_lbl.bold = True
        r_inst_lbl.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
        p_inst.add_run(item['instruction'])

        # Placeholder Box for image
        box_table = doc.add_table(rows=1, cols=1)
        box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        box_cell = box_table.rows[0].cells[0]
        box_cell.width = Inches(6.5)
        set_cell_background(box_cell, 'F8FAFC')
        set_cell_margins(box_cell, 400, 400, 200, 200)

        box_p = box_cell.paragraphs[0]
        box_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        box_p.paragraph_format.space_before = Pt(20)
        box_p.paragraph_format.space_after = Pt(20)
        box_r = box_p.add_run(f"[ CHÈN ẢNH VÀO ĐÂY: {item['title']} ]\n(Dán ảnh chụp màn hình thực tế của bạn tại vị trí này)")
        box_r.bold = True
        box_r.font.size = Pt(11)
        box_r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # Exact Caption Required
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(16)
        r_cap_tag = p_cap.add_run("Caption: ")
        r_cap_tag.bold = True
        r_cap_tag.font.size = Pt(10)
        r_cap_tag.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
        r_cap_text = p_cap.add_run(item['caption'])
        r_cap_text.italic = True
        r_cap_text.font.size = Pt(10)
        r_cap_text.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Section 3: Hướng dẫn đóng gói nộp bài
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    h3_run = h3.add_run("3. QUY ĐỊNH ĐÓNG GÓI VÀ NỘP BÀI (SUBMISSION GUIDELINES)")
    h3_run.font.size = Pt(14)
    h3_run.font.bold = True
    h3_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_sub = doc.add_paragraph(
        "Theo Mục 4 của Đề thi, sinh viên tạo một thư mục mang tên <StudentCode> chứa 2 tệp tin sau và nộp:\n\n"
        "1. Tệp Word: <StudentCode>.docx — chính là tệp này sau khi bạn đã điền thông tin và chèn đầy đủ 16 ảnh chụp màn hình.\n"
        "2. Tệp nén mã nguồn: <StudentCode>.zip — nén toàn bộ mã nguồn sạch của dự án, bao gồm:\n"
        "   - package.json ở ngay cấp thư mục gốc của file zip (không nằm trong thư mục con lồng nhau).\n"
        "   - package-lock.json (hoặc bun.lock / yarn.lock).\n"
        "   - Thư mục src/ (chứa screens, components, contexts, hooks, services, utils, constants, types).\n"
        "   - Tệp firestore.rules.\n"
        "   - Tệp .env.example (chứa tên biến cấu hình Firebase không có giá trị bí mật).\n"
        "   - LƯU Ý ĐẶC BIỆT: TUYỆT ĐỐI KHÔNG nén các thư mục node_modules, .expo, .git hoặc tệp .env thật có chứa giá trị bí mật."
    )
    p_sub.paragraph_format.space_after = Pt(12)

    output_path = "D:\\MMA301_PE\\QE200012_MMA301_REPORT.docx"
    doc.save(output_path)
    print(f"Document created successfully at: {output_path}")

if __name__ == '__main__':
    create_report()
